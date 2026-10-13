"""Extract text from an uploaded PDF/DOCX CV and turn it into structured profile data."""
import io
import json
import os
import re
from datetime import date

import requests


def extract_text(uploaded):
    name = uploaded.name.lower()
    data = uploaded.read()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    raise ValueError("Only .pdf and .docx files are supported.")


PROMPT = """Extract this CV into JSON only (no prose, no code fences) with exactly these keys:
full_name, headline, phone, address, summary, linkedin_url, github_url, portfolio_url,
education: [{institute, degree, start_year(int), end_year(int|null), description}],
experience: [{company, role, start_date(YYYY-MM-DD), end_date(YYYY-MM-DD|null), is_current(bool), description}],
skills: [string], projects: [{title, description, link}],
certifications: [{title, issuer, year(int|null)}].
Use "" or [] when missing. Never invent data.

CV:
"""


GEMINI_MODELS = ["gemini-flash-latest", "gemini-2.5-flash-lite", "gemini-2.5-flash", "gemini-2.0-flash", "gemini-flash-lite-latest"]


def parse_with_ai(text):
    key = os.environ.get("GEMINI_API_KEY")
    if not key:
        return None
    first = os.environ.get("GEMINI_MODEL")  # optional override in .env
    models = ([first] if first else []) + [m for m in GEMINI_MODELS if m != first]
    last_err = None
    for model in models:
        r = requests.post(
            f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent",
            headers={"x-goog-api-key": key, "content-type": "application/json"},
            json={
                "contents": [{"parts": [{"text": PROMPT + text[:20000]}]}],
                "generationConfig": {"responseMimeType": "application/json", "temperature": 0},
            },
            timeout=60,
        )
        if r.status_code in (404, 429, 503):  # model unavailable / quota hit / overloaded -> try the next one
            last_err = f"{model}: HTTP {r.status_code} {r.text[:150]}"
            continue
        r.raise_for_status()
        raw = r.json()["candidates"][0]["content"]["parts"][0]["text"].strip()
        raw = re.sub(r"^```(?:json)?|```$", "", raw, flags=re.M).strip()
        return json.loads(raw)
    raise RuntimeError(f"No Gemini model worked. Last: {last_err}")


_MON = {m: i for i, m in enumerate(["jan","feb","mar","apr","may","jun","jul","aug","sep","oct","nov","dec"], 1)}
_SECTIONS = [
    ("summary", r"summary|profile|objective|about"),
    ("skills", r"skills|technologies|tech stack"),
    ("experience", r"experience|employment|work history"),
    ("projects", r"projects"),
    ("education", r"education|academic"),
    ("certifications", r"certif|achievements|awards"),
    ("languages", r"^languages$"),
]
_DASH = r"\s[\u2014\u2013-]\s"
_DATES = r"((?:[A-Za-z]{3,9}\.?\s+)?\d{4})\s*[\u2013\u2014-]\s*(Present|Current|Now|(?:[A-Za-z]{3,9}\.?\s+)?\d{4})"
_BULLET = re.compile(r"^[^\w\s(\"'+\[]\s*")  # any leading symbol: \u2022, \x7f, -, *, private-use glyphs...


def _to_date(txt):
    m = re.match(r"(?:([A-Za-z]{3})[a-z]*\.?\s+)?(\d{4})", (txt or "").strip())
    if not m:
        return None
    return f"{m.group(2)}-{_MON.get((m.group(1) or 'jan').lower(), 1):02d}-01"


def _section_of(line):
    t = re.sub(r"[^A-Za-z& ]", "", line).strip()
    if not t or len(t) > 40 or not (line.isupper() or line.istitle() and len(t.split()) <= 3):
        return None
    for name, pat in _SECTIONS:
        if name == "languages" and not line.isupper():
            continue  # "Languages" inside a skills list is a label, not a section
        if re.search(pat, t.lower()):
            return name
    return None


def parse_fallback(text):
    """No AI available: rule-based parser for common CV layouts (best effort)."""
    lines = [re.sub(r"\s+", " ", l).strip() for l in text.splitlines()]
    lines = [l for l in lines if l]
    out = {"full_name": lines[0][:200] if lines else "", "headline": "", "phone": "", "address": "", "summary": "",
           "linkedin_url": "", "github_url": "", "portfolio_url": "",
           "education": [], "experience": [], "skills": [], "projects": [], "certifications": []}

    # --- header (everything before the first section heading) ---
    first = next((i for i, l in enumerate(lines) if _section_of(l)), len(lines))
    head = lines[1:first]
    if head and "@" not in head[0] and not re.search(r"\d{5,}", head[0]) and len(head[0]) < 80:
        out["headline"] = head.pop(0)
    blob = " | ".join(head)
    m = re.search(r"\+?\d[\d\s\-().]{8,}\d", blob)
    out["phone"] = m.group(0).strip() if m else ""
    out["linkedin_url"] = (re.search(r"(?:https?://)?(?:www\.)?linkedin\.com/[^\s|]+", blob, re.I) or [""])[0]
    out["github_url"] = (re.search(r"(?:https?://)?(?:www\.)?github\.com/[^\s|]+", blob, re.I) or [""])[0]
    for part in re.split(r"\s*\|\s*", blob):
        if "," in part and not re.search(r"@|\d{4,}|\.com|\.dev|/", part):
            out["address"] = part.strip()
            break

    # --- split into sections ---
    secs, cur = {}, None
    for l in lines[first:]:
        name = _section_of(l)
        if name:
            cur = name
            secs[cur] = []
        elif cur:
            secs[cur].append(l)

    out["summary"] = " ".join(secs.get("summary", []))

    cat = re.compile(r"^(programming )?(languages?|frontend|front-end|backend|back-end|databases?|devops(?: ?/ ?tools)?|tools|cloud|frameworks?|libraries|other|mobile|testing)\s*:?\s*$", re.I)
    for l in secs.get("skills", []):
        l = re.sub(r"^[^:]{2,25}:\s*", "", l) if ":" in l[:26] else l
        if cat.match(l):
            continue
        for sk in re.split(r"[,\u2022|;](?![^(]*\))", l):
            sk = sk.strip()
            if 1 < len(sk) < 40 and sk.lower() not in {x.lower() for x in out["skills"]}:
                out["skills"].append(sk)

    # --- experience / education: "Title - Org   dates" (dates may sit on the next line) ---
    def entries(sec_lines, kind):
        items, item = [], None
        for i, l in enumerate(sec_lines):
            body = _BULLET.sub("", l) if _BULLET.match(l) else None
            dm = re.search(_DATES, l)
            date_only = bool(re.fullmatch(r"\s*" + _DATES + r".*", l))
            is_head = (not body) and (not date_only) and re.search(_DASH, re.sub(_DATES, "", l))
            nxt_dm = i + 1 < len(sec_lines) and re.fullmatch(r"\s*" + _DATES + r".*", sec_lines[i + 1])
            if is_head and (dm or nxt_dm):
                title_part = re.sub(_DATES + r".*$", "", l).strip() if dm else l
                bits = [x.strip() for x in re.split(_DASH, title_part, maxsplit=1)]
                a, b = bits[0], (bits[1] if len(bits) > 1 else "")
                item = {"a": a, "b": b, "dates": dm.groups() if dm else None, "desc": []}
                items.append(item)
            elif item is not None and item["dates"] is None and date_only:
                item["dates"] = dm.groups()
            elif item is not None and body:
                item["desc"].append(body)
            elif item is not None and item["desc"] and not date_only:
                item["desc"][-1] += " " + l  # wrapped continuation of the previous bullet
        return items

    for e in entries(secs.get("experience", []), "exp"):
        start = _to_date(e["dates"][0]) if e["dates"] else None
        if not start:
            continue
        cur_job = bool(re.match(r"present|current|now", e["dates"][1], re.I))
        out["experience"].append({"role": e["a"], "company": e["b"], "start_date": start,
                                  "end_date": None if cur_job else _to_date(e["dates"][1]),
                                  "is_current": cur_job, "description": "\n".join(e["desc"])})
    for e in entries(secs.get("education", []), "edu"):
        if not e["dates"]:
            continue
        end = e["dates"][1]
        out["education"].append({"degree": e["a"], "institute": e["b"],
                                 "start_year": int(re.search(r"\d{4}", e["dates"][0]).group(0)),
                                 "end_year": None if re.match(r"present|current|now", end, re.I) else int(re.search(r"\d{4}", end).group(0)),
                                 "description": "\n".join(e["desc"])})

    # --- projects: "Title - subtitle  tech" then bullets ---
    proj = None
    for l in secs.get("projects", []):
        if _BULLET.match(l):
            if proj is not None:
                proj["description"] += ("\n" if proj["description"] else "") + _BULLET.sub("", l)
                proj["_b"] = True
        elif proj is not None and not re.search(_DASH, l) and (not proj.get("_b") or l[:1].islower()):
            if proj.get("_b"):
                proj["description"] += " " + l  # wrapped bullet
            else:
                proj["description"] += ("\n" if proj["description"] else "") + "Tech: " + l  # tech-stack line
        else:
            parts = re.split(_DASH, l, maxsplit=1)
            proj = {"title": parts[0].strip(), "description": parts[1].strip() if len(parts) > 1 else "", "link": ""}
            out["projects"].append(proj)

    for l in secs.get("certifications", []):
        t = _BULLET.sub("", l)
        y = re.search(r"\(?\b((?:19|20)\d{2})\b\)?", t)
        title = re.sub(r"\s*\(\s*(?:19|20)\d{2}\s*\)", "", t).strip()
        if title:
            out["certifications"].append({"title": title, "issuer": "", "year": int(y.group(1)) if y else None})
    for p in out["projects"]:
        p.pop("_b", None)
    return out


def _d(v):
    try:
        return date.fromisoformat(v) if v else None
    except ValueError:
        return None


def _i(v):
    try:
        return int(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _norm(x):
    return re.sub(r"[^a-z0-9]+", "", (x or "").lower())


def _same(a, b):
    """Loose text equality: 'Brightwave Technologies' == 'Brightwave'."""
    a, b = _norm(a), _norm(b)
    return bool(a and b) and (a == b or (min(len(a), len(b)) >= 5 and (a in b or b in a)))


def apply_to_profile(profile, d):
    """Fill the profile from parsed CV data WITHOUT creating duplicates.
    Returns {"added": {section: n}, "skipped": n}."""
    from .models import Certification, Education, Experience, Project, Skill

    added = {"experience": 0, "education": 0, "skills": 0, "projects": 0, "certifications": 0}
    skipped = 0

    for f in ("full_name", "headline", "phone", "address", "summary", "linkedin_url", "github_url", "portfolio_url"):
        if d.get(f) and not getattr(profile, f):  # never overwrite what the user already typed
            setattr(profile, f, str(d[f])[:255 if f == "address" else 200 if f == "full_name" else 150 if f == "headline" else 30 if f == "phone" else 100000])
    profile.save()

    exp = list(profile.experience.all())
    for e in d.get("experience") or []:
        sd = _d(e.get("start_date"))
        if not (e.get("company") and sd):
            continue
        if any(_same(x.company, e["company"]) and (x.start_date == sd or _same(x.role, e.get("role"))) for x in exp):
            skipped += 1
            continue
        exp.append(Experience.objects.create(profile=profile, company=e["company"][:200], role=(e.get("role") or "")[:200],
                                             start_date=sd, end_date=_d(e.get("end_date")),
                                             is_current=bool(e.get("is_current")), description=e.get("description") or ""))
        added["experience"] += 1

    edu = list(profile.education.all())
    for e in d.get("education") or []:
        sy = _i(e.get("start_year"))
        if not (e.get("institute") and sy):
            continue
        if any(_same(x.institute, e["institute"]) and (x.start_year == sy or _same(x.degree, e.get("degree"))) for x in edu):
            skipped += 1
            continue
        edu.append(Education.objects.create(profile=profile, institute=e["institute"][:200], degree=(e.get("degree") or "")[:200],
                                            start_year=sy, end_year=_i(e.get("end_year")), description=e.get("description") or ""))
        added["education"] += 1

    have = {sk.name.lower() for sk in profile.skills.all()}
    for sk in d.get("skills") or []:
        if not (isinstance(sk, str) and sk.strip()):
            continue
        if sk.strip().lower() in have:
            skipped += 1
            continue
        Skill.objects.create(profile=profile, name=sk.strip()[:100])
        have.add(sk.strip().lower())
        added["skills"] += 1

    projs = list(profile.projects.all())
    for p in d.get("projects") or []:
        if not p.get("title"):
            continue
        if any(_same(x.title, p["title"]) for x in projs):
            skipped += 1
            continue
        projs.append(Project.objects.create(profile=profile, title=p["title"][:200], description=p.get("description") or "",
                                            link=p.get("link") if str(p.get("link", "")).startswith("http") else ""))
        added["projects"] += 1

    certs = list(profile.certifications.all())
    for c in d.get("certifications") or []:
        if not c.get("title"):
            continue
        if any(_same(x.title, c["title"]) for x in certs):
            skipped += 1
            continue
        certs.append(Certification.objects.create(profile=profile, title=c["title"][:200], issuer=(c.get("issuer") or "")[:200], year=_i(c.get("year"))))
        added["certifications"] += 1

    return {"added": added, "skipped": skipped}
